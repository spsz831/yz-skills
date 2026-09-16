# -*- coding: utf-8 -*-
"""审计影视收藏表，发现重复、缺字段和链接问题。"""
import argparse
import json
import re
import urllib.error
import urllib.request
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


def norm(value):
    if value is None:
        return ""
    return re.sub(r"\s+", "", str(value).replace("《", "").replace("》", "")).lower()


def url_status(url, timeout):
    if not isinstance(url, str) or not re.match(r"^https?://", url, re.I):
        return {"ok": False, "status": "invalid_format"}
    request = urllib.request.Request(url, method="HEAD", headers={"User-Agent": "Mozilla/5.0"})
    try:
        with urllib.request.urlopen(request, timeout=timeout) as response:
            return {"ok": 200 <= response.status < 400, "status": response.status}
    except urllib.error.HTTPError as exc:
        if exc.code in (403, 405):
            try:
                request = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0", "Range": "bytes=0-0"})
                with urllib.request.urlopen(request, timeout=timeout) as response:
                    return {"ok": 200 <= response.status < 400, "status": response.status}
            except Exception as retry_exc:
                return {"ok": False, "status": f"{type(retry_exc).__name__}: {retry_exc}"}
        return {"ok": False, "status": exc.code}
    except Exception as exc:
        return {"ok": False, "status": f"{type(exc).__name__}: {exc}"}


def audit_sheet(ws, check_links, timeout):
    headers = {str(cell.value).strip(): cell.column for cell in ws[1] if cell.value not in (None, "")}
    rows = []
    duplicate_map = defaultdict(list)
    for row_no in range(2, ws.max_row + 1):
        cells = list(ws[row_no])
        def get(name):
            if name not in headers:
                return None
            cell = cells[headers[name] - 1]
            if name in ("网盘链接", "豆瓣链接") and cell.hyperlink:
                return cell.hyperlink.target
            return cell.value
        name = get("电影名称")
        if name in (None, ""):
            continue
        item = {
            "sheet": ws.title,
            "row": row_no,
            "name": str(name),
            "media_type": get("媒体类型"),
            "year": get("年份"),
            "rating": get("豆瓣评分"),
            "episodes": get("集数"),
            "seasons": get("季数"),
            "link": get("网盘链接"),
            "douban_url": get("豆瓣链接"),
        }
        rows.append(item)
        duplicate_map[norm(name)].append(item)

    duplicate_titles = [items for items in duplicate_map.values() if len(items) > 1]
    missing_rating = [item for item in rows if item["rating"] in (None, "")]
    missing_link = [item for item in rows if item["link"] in (None, "")]
    missing_douban_url = [item for item in rows if item["douban_url"] in (None, "")]
    is_series = ws.title == "电视剧收藏"
    missing_episodes = [item for item in rows if is_series and item["episodes"] in (None, "")]
    missing_seasons = [item for item in rows if is_series and item["seasons"] in (None, "")]
    invalid_links = []
    dead_links = []
    if check_links:
        for item in rows:
            for field in ("link", "douban_url"):
                value = item[field]
                if value in (None, ""):
                    continue
                result = url_status(value, timeout)
                if not result["ok"]:
                    dead_links.append({**item, "field": field, "url": value, "status": result["status"]})
                elif not re.match(r"^https?://", str(value), re.I):
                    invalid_links.append({**item, "field": field, "url": value})

    return {
        "sheet": ws.title,
        "rows": len(rows),
        "duplicate_titles": duplicate_titles,
        "missing_rating": missing_rating,
        "missing_link": missing_link,
        "missing_douban_url": missing_douban_url,
        "missing_episodes": missing_episodes,
        "missing_seasons": missing_seasons,
        "invalid_links": invalid_links,
        "dead_links": dead_links,
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--xlsx", required=True)
    parser.add_argument("--check-links", action="store_true", help="实际访问网盘/豆瓣链接，检查 HTTP 响应")
    parser.add_argument("--timeout", type=float, default=8)
    args = parser.parse_args()
    workbook = load_workbook(Path(args.xlsx), read_only=False, data_only=False)
    result = {
        "ok": True,
        "xlsx": str(Path(args.xlsx)),
        "check_links": args.check_links,
        "sheets": [audit_sheet(workbook[name], args.check_links, args.timeout) for name in workbook.sheetnames if name in ("电影收藏", "电视剧收藏")],
    }
    result["summary"] = {
        key: sum(len(sheet[key]) for sheet in result["sheets"])
        for key in ("duplicate_titles", "missing_rating", "missing_link", "missing_douban_url", "missing_episodes", "missing_seasons", "invalid_links", "dead_links")
    }
    print(json.dumps(result, ensure_ascii=False, indent=2, default=str))


if __name__ == "__main__":
    main()
